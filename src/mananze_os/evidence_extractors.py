
"""Governed document extractors for Mananze OS evidence ingestion."""

from __future__ import annotations

import io
import json
from typing import Any

from .evidence_content import (
    EvidenceContent,
    EvidenceContentExtractor,
    EvidenceExtractionRequest,
    EvidenceContentIngestion,
)


class PDFEvidenceExtractor:
    """Extract text from PDF evidence while preserving page boundaries."""

    extractor_id = "mananze.pdf.pypdf"
    supported_content_types = frozenset({"application/pdf"})

    def extract(self, request: EvidenceExtractionRequest) -> EvidenceContent:
        if request.content_type != "application/pdf":
            raise ValueError("PDF extractor requires application/pdf")

        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(request.raw_content))
        pages: list[dict[str, Any]] = []

        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append(
                {
                    "page": index,
                    "text": text,
                }
            )

        payload = {
            "format": "pdf",
            "page_count": len(pages),
            "pages": pages,
        }

        return EvidenceContent(
            evidence_id=request.evidence.evidence_id,
            tenant_id=request.tenant_id,
            content_type=request.content_type,
            content=json.dumps(payload, ensure_ascii=False, sort_keys=True),
            content_hash=_sha256(request.raw_content),
            extractor_id=self.extractor_id,
        )


class DOCXEvidenceExtractor:
    """Extract paragraphs and tables from DOCX evidence."""

    extractor_id = "mananze.docx.python-docx"
    supported_content_types = frozenset(
        {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        }
    )

    def extract(self, request: EvidenceExtractionRequest) -> EvidenceContent:
        if request.content_type not in self.supported_content_types:
            raise ValueError("DOCX extractor requires the DOCX content type")

        from docx import Document

        document = Document(io.BytesIO(request.raw_content))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        tables: list[list[list[str]]] = []

        for table in document.tables:
            rows: list[list[str]] = []
            for row in table.rows:
                rows.append([cell.text for cell in row.cells])
            tables.append(rows)

        payload = {
            "format": "docx",
            "paragraphs": paragraphs,
            "tables": tables,
        }

        return EvidenceContent(
            evidence_id=request.evidence.evidence_id,
            tenant_id=request.tenant_id,
            content_type=request.content_type,
            content=json.dumps(payload, ensure_ascii=False, sort_keys=True),
            content_hash=_sha256(request.raw_content),
            extractor_id=self.extractor_id,
        )


class XLSXEvidenceExtractor:
    """Extract workbook sheets and cell values from XLSX evidence."""

    extractor_id = "mananze.xlsx.openpyxl"
    supported_content_types = frozenset(
        {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        }
    )

    def extract(self, request: EvidenceExtractionRequest) -> EvidenceContent:
        if request.content_type not in self.supported_content_types:
            raise ValueError("XLSX extractor requires the XLSX content type")

        from openpyxl import load_workbook

        workbook = load_workbook(
            io.BytesIO(request.raw_content),
            read_only=True,
            data_only=False,
        )

        sheets: list[dict[str, Any]] = []

        for worksheet in workbook.worksheets:
            rows: list[list[Any]] = []

            for row in worksheet.iter_rows(values_only=True):
                rows.append(list(row))

            sheets.append(
                {
                    "name": worksheet.title,
                    "rows": rows,
                }
            )

        workbook.close()

        payload = {
            "format": "xlsx",
            "sheet_count": len(sheets),
            "sheets": sheets,
        }

        return EvidenceContent(
            evidence_id=request.evidence.evidence_id,
            tenant_id=request.tenant_id,
            content_type=request.content_type,
            content=json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str),
            content_hash=_sha256(request.raw_content),
            extractor_id=self.extractor_id,
        )


def _sha256(content: bytes) -> str:
    import hashlib

    return hashlib.sha256(content).hexdigest()


PDF_CONTENT_TYPE = "application/pdf"
DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


__all__ = [
    "PDFEvidenceExtractor",
    "DOCXEvidenceExtractor",
    "XLSXEvidenceExtractor",
    "PDF_CONTENT_TYPE",
    "DOCX_CONTENT_TYPE",
    "XLSX_CONTENT_TYPE",
]
